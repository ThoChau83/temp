"""Sinh các tài liệu tri thức dạng bảng và PDF từ master_data.py + policy.yaml.

    python tools/build_docs.py

Đầu ra (thư mục knowledge/):
- 22-bang-gia-hoa-cuoi-su-kien-2026.pdf
- 23-dieu-khoan-hop-dong-hoa-su-kien.pdf   (nguồn: tools/pdf_src/23-...md)
- 24-catalogue-hoa-20-10-2026.pdf
- 25-bang-gia-san-pham.csv
- 26-ma-khuyen-mai-2026.csv
- 27-phi-giao-hang-khu-vuc.csv
- 28-lich-mua-hoa.csv
- 29-bang-gia-goi-dinh-ky-va-su-kien.xlsx
"""
from __future__ import annotations

import csv
import re
import sys
from datetime import datetime
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

TOOLS = Path(__file__).resolve().parent
ROOT = TOOLS.parent
OUT = ROOT / "knowledge"
sys.path.insert(0, str(TOOLS))
import master_data as M  # noqa: E402

POLICY = yaml.safe_load((ROOT / "policy.yaml").read_text(encoding="utf-8"))
ANCHOR = datetime.fromisoformat(M.ANCHOR_NOW).date().isoformat()
BRAND = colors.HexColor("#9F1D20")
SOFT = colors.HexColor("#F7E3DA")


def vnd(n: int | None) -> str:
    return "" if n is None else f"{int(n):,}".replace(",", ".") + "đ"


# ============================================================ CSV
def write_csv(name: str, header: list[str], rows: list[list]) -> None:
    with open(OUT / name, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    print(f"  {name}: {len(rows)} dòng")


def build_csvs() -> None:
    write_csv("25-bang-gia-san-pham.csv",
              ["sku", "ten_san_pham", "danh_muc", "gia_vnd", "don_vi", "hoa_chinh", "tong_mau", "kich_thuoc",
               "thoi_gian_chuan_bi_gio", "cam_ket_tuoi_ngay", "theo_mua", "ghi_chu_mua", "mo_ta"],
              [[p["sku"], p["name"], M.CATEGORIES[p["category"]], p["price"], p["unit"], p["flowers"], p["tone"],
                p["size"], max(p["lead"], POLICY["ordering"]["min_lead_time_hours"]), p["fresh"] if p["fresh"] else "",
                "có" if p.get("seasonal") else "không", p.get("season_note") or "", p["desc"]] for p in M.PRODUCTS])

    def status(c: dict) -> str:
        if ANCHOR < c["valid_from"]:
            return "sắp diễn ra"
        return "hết hạn" if ANCHOR > c["valid_to"] else "đang áp dụng"

    write_csv("26-ma-khuyen-mai-2026.csv",
              ["ma", "mo_ta", "loai", "gia_tri", "giam_toi_da_vnd", "don_toi_thieu_vnd", "danh_muc_ap_dung",
               "tu_ngay", "den_ngay", "so_lan_moi_khach", "tong_luot", f"trang_thai_ngay_{ANCHOR}"],
              [[c["code"], c["description"], "phần trăm" if c["discount_type"] == "percent" else "số tiền",
                f"{c['value']}%" if c["discount_type"] == "percent" else c["value"], c["max_discount"], c["min_order"],
                ", ".join(M.CATEGORIES[x] for x in c["applicable_categories"].split(",")) if c["applicable_categories"] else "Tất cả",
                c["valid_from"], c["valid_to"], c["per_customer_limit"] or "", c["usage_limit"] or "không giới hạn",
                status(c)] for c in M.DISCOUNT_CODES])

    express = POLICY["ordering"]["express"]
    branch_name = {b["id"]: b["name"].split(" (")[0] for b in M.BRANCHES}
    rows = [[z["id"], z["name"], z["areas"], z["fee"], z["free_ship_min_order"] or "không áp dụng",
             z["allowed_slots"].replace(",", ", "), "có" if z["express_available"] else "không",
             z["min_days_ahead"], branch_name[z["default_branch_id"]]] for z in M.ZONES]
    rows.append(["NGOAI_VUNG", "Ngoài vùng giao", "Vị Thanh, Ngã Bảy, Sóc Trăng, tỉnh khác", "báo giá riêng", "",
                 "", "không", "", "Trưởng ca duyệt (APR-06)"])
    rows.append(["HOA_TOC", "Phụ phí giao hỏa tốc 2 giờ", ", ".join(express["zones"]), express["fee"], "",
                 f"{express['start']}–{express['end']}", "có", 0, "Không áp dụng ngày cao điểm"])
    write_csv("27-phi-giao-hang-khu-vuc.csv",
              ["ma_khu_vuc", "ten_khu_vuc", "khu_vuc_bao_gom", "phi_ship_vnd", "mien_phi_tu_vnd", "khung_gio",
               "giao_hoa_toc", "dat_truoc_toi_thieu_ngay", "chi_nhanh_phu_trach"], rows)

    legend = {"N": "nhiều", "C": "có", "I": "ít", "K": "không"}
    write_csv("28-lich-mua-hoa.csv",
              ["loai_hoa", "nguon"] + [f"T{i}" for i in range(1, 13)] + ["dip_gia_cao", "sku_lien_quan"],
              [[name, src] + [legend[ch] for ch in months] + [peak, skus]
               for name, src, months, peak, skus in M.SEASONAL_CALENDAR])


# ============================================================ XLSX
def build_xlsx() -> None:
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="9F1D20")

    def sheet(ws, header, rows, widths, money_cols=()):
        ws.append(header)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for col in money_cols:
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                row[0].number_format = "#,##0"
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"

    disc = POLICY["subscription"]["term_discount_percent"]
    ws = wb.active
    ws.title = "Goi_dinh_ky"
    rows = []
    for p in M.SUBSCRIPTION_PLANS:
        row = [p["id"], p["name"], {"weekly": "Mỗi tuần", "biweekly": "2 tuần/lần",
                                    "lunar_1_15": "Trước mùng 1 và rằm âm lịch"}[p["frequency"]],
               p["description"], p["price_per_delivery"], p["retail_equivalent"]]
        for months in (1, 3, 6):
            n = p["deliveries_per_month"] * months
            pct = disc[str(months)]
            row += [n, int(p["price_per_delivery"] * n * (100 - pct) / 100 // 1000 * 1000)]
        rows.append(row)
    sheet(ws, ["Mã gói", "Tên gói", "Tần suất", "Mô tả", "Giá/lần", "Giá lẻ tương đương",
               "1 tháng – số lần", "1 tháng – tổng", "3 tháng (−5%) – số lần", "3 tháng – tổng",
               "6 tháng (−10%) – số lần", "6 tháng – tổng"], rows,
          [14, 24, 24, 50, 12, 14, 10, 14, 12, 14, 12, 14], money_cols=(5, 6, 8, 10, 12))

    ws = wb.create_sheet("Goi_su_kien")
    sheet(ws, ["Mã gói", "Loại sự kiện", "Tên gói", "Giá trọn gói (VND)", "Quy mô", "Hạng mục bao gồm"],
          [[p["code"], M.EVENT_TYPES[p["event_type"]], p["name"], p["price"], p["scale"], "\n".join(p["items"])]
           for p in M.EVENT_PACKAGES], [16, 24, 28, 16, 26, 60], money_cols=(4,))

    ws = wb.create_sheet("Hang_muc_them")
    sheet(ws, ["Hạng mục", "Đơn giá (VND)", "Đơn vị"], [list(a) for a in M.EVENT_ADDONS], [36, 16, 10],
          money_cols=(2,))

    ws = wb.create_sheet("Phu_phi_va_coc")
    con = POLICY["consultation"]
    rows = [[name, value] for name, value in M.EVENT_SURCHARGES]
    rows += [["Đặt cọc khi ký hợp đồng", f"{con['contract']['deposit_percent']}% giá trị hợp đồng"],
             ["Thanh toán còn lại", f"Trước ngày sự kiện {con['contract']['balance_due_days_before_event']} ngày"],
             ["Đổi ngày sự kiện miễn phí", f"{con['contract']['free_event_date_change']['times']} lần, báo trước "
                                           f"≥ {con['contract']['free_event_date_change']['min_days_before']} ngày"]]
    for t in con["deposit_refund_tiers"]:
        rows.append([f"Huỷ khi còn ≥ {t['min_days_before_event']} ngày trước sự kiện",
                     f"Hoàn {t['refund_percent']}% tiền cọc (quản lý studio duyệt)"])
    sheet(ws, ["Khoản mục", "Quy định"], rows, [52, 52])

    wb.properties.title = "Bảng giá gói định kỳ và sự kiện – Mộc Miên Florist"
    wb.properties.creator = "Mộc Miên Florist"
    path = OUT / "29-bang-gia-goi-dinh-ky-va-su-kien.xlsx"
    wb.save(path)
    print(f"  {path.name}: {len(wb.sheetnames)} sheet")


# ============================================================ PDF
def register_fonts() -> tuple[str, str]:
    candidates = [
        ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
    ]
    for regular, bold in candidates:
        if Path(regular).exists() and Path(bold).exists():
            pdfmetrics.registerFont(TTFont("VN", regular))
            pdfmetrics.registerFont(TTFont("VN-Bold", bold))
            pdfmetrics.registerFontFamily("VN", normal="VN", bold="VN-Bold", italic="VN", boldItalic="VN-Bold")
            return "VN", "VN-Bold"
    raise SystemExit("Không tìm thấy font Unicode (Arial/DejaVu) để in tiếng Việt.")


FONT, FONT_BOLD = "Helvetica", "Helvetica-Bold"
STYLES: dict[str, ParagraphStyle] = {}


def init_styles() -> None:
    global FONT, FONT_BOLD
    FONT, FONT_BOLD = register_fonts()
    base = dict(fontName=FONT, textColor=colors.HexColor("#2B2B2B"))
    STYLES.update({
        "h1": ParagraphStyle("h1", fontName=FONT_BOLD, fontSize=19, leading=24, textColor=BRAND, spaceAfter=8),
        "h2": ParagraphStyle("h2", fontName=FONT_BOLD, fontSize=13.5, leading=18, textColor=BRAND, spaceBefore=12, spaceAfter=5),
        "h3": ParagraphStyle("h3", fontName=FONT_BOLD, fontSize=11.5, leading=15, spaceBefore=8, spaceAfter=3,
                             textColor=colors.HexColor("#5A2A1A")),
        "p": ParagraphStyle("p", fontSize=10.2, leading=14.5, spaceAfter=5, alignment=TA_LEFT, **base),
        "cell": ParagraphStyle("cell", fontSize=9.2, leading=12, **base),
        "cellh": ParagraphStyle("cellh", fontName=FONT_BOLD, fontSize=9.2, leading=12, textColor=colors.white),
        "note": ParagraphStyle("note", fontSize=9.6, leading=13.5, **base),
    })


def inline(text: str) -> str:
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    return re.sub(r"`(.+?)`", r"<b>\1</b>", text)


def md_to_flowables(md: str) -> list:
    flow, para, lines = [], [], md.splitlines()
    i = 0

    def flush():
        if para:
            flow.append(Paragraph(inline(" ".join(para)), STYLES["p"]))
            para.clear()

    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            flush()
            i += 1
        elif line.startswith("### "):
            flush()
            flow.append(Paragraph(inline(line[4:]), STYLES["h3"]))
            i += 1
        elif line.startswith("## "):
            flush()
            flow.append(Paragraph(inline(line[3:]), STYLES["h2"]))
            i += 1
        elif line.startswith("# "):
            flush()
            flow.append(Paragraph(inline(line[2:]), STYLES["h1"]))
            i += 1
        elif line.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            data = [[Paragraph(inline(c), STYLES["cellh" if r == 0 else "cell"]) for c in row]
                    for r, row in enumerate(rows)]
            width = A4[0] - 36 * mm
            table = Table(data, colWidths=[width / len(rows[0])] * len(rows[0]), repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), BRAND),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FDF4EF")]),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E0C9BE")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            flow += [table, Spacer(1, 6)]
        elif re.match(r"^(- |\d+\. )", line):
            flush()
            numbered = bool(re.match(r"^\d+\. ", line))
            items = []
            while i < len(lines) and re.match(r"^(- |\d+\. )", lines[i]):
                items.append(ListItem(Paragraph(inline(re.sub(r"^(- |\d+\. )", "", lines[i])), STYLES["p"]),
                                      leftIndent=14))
                i += 1
            flow.append(ListFlowable(items, bulletType="1" if numbered else "bullet", bulletFontName=FONT,
                                     bulletFontSize=9, leftIndent=14, start="1" if numbered else "•"))
        elif line.startswith("> "):
            flush()
            note = []
            while i < len(lines) and lines[i].startswith("> "):
                note.append(lines[i][2:])
                i += 1
            box = Table([[Paragraph(inline(" ".join(note)), STYLES["note"])]], colWidths=[A4[0] - 36 * mm])
            box.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), SOFT),
                                     ("LINEBEFORE", (0, 0), (0, -1), 3, BRAND),
                                     ("LEFTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 6),
                                     ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
            flow += [box, Spacer(1, 6)]
        else:
            para.append(line.strip())
            i += 1
    flush()
    return flow


def render_pdf(name: str, title: str, md: str) -> None:
    path = OUT / name

    def decorate(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(BRAND)
        canvas.rect(0, A4[1] - 12 * mm, A4[0], 12 * mm, fill=1, stroke=0)
        canvas.setFillColor(colors.white)
        canvas.setFont(FONT_BOLD, 10)
        canvas.drawString(18 * mm, A4[1] - 8 * mm, "MỘC MIÊN FLORIST")
        canvas.setFont(FONT, 9)
        canvas.drawRightString(A4[0] - 18 * mm, A4[1] - 8 * mm, title)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.setFont(FONT, 8)
        canvas.drawString(18 * mm, 10 * mm, "86 Trần Văn Hoài, phường Ninh Kiều, TP. Cần Thơ · Hotline 0292 3 868 186 · mocmienflorist.vn")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Trang {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=20 * mm,
                            bottomMargin=18 * mm, title=title, author="Mộc Miên Florist", subject=title,
                            creator="tools/build_docs.py")
    doc.build(md_to_flowables(md), onFirstPage=decorate, onLaterPages=decorate)
    print(f"  {name}")


def wedding_price_md() -> str:
    lines = ["# Bảng giá hoa cưới & sự kiện 2026",
             "Áp dụng từ 01/10/2026 đến 31/12/2026 tại Studio Mộc Miên (tầng 2, 86 Trần Văn Hoài, phường Ninh Kiều). "
             "Giá đã gồm thuế GTGT, vận chuyển và setup trong khu vực Z1–Z2. Tư vấn tại studio và online **miễn phí**.",
             ""]
    for etype, label in M.EVENT_TYPES.items():
        lines += [f"## {label}", ""]
        for p in [x for x in M.EVENT_PACKAGES if x["event_type"] == etype]:
            lines += [f"### {p['name']} – {vnd(p['price'])}", f"Mã gói: {p['code']} · Quy mô: {p['scale']}.", ""]
            lines += [f"- {item}" for item in p["items"]]
            lines.append("")
    lines += ["## Hạng mục thêm", "", "| Hạng mục | Đơn giá | Đơn vị |", "|---|---|---|"]
    lines += [f"| {n} | {vnd(price)} | {unit} |" for n, price, unit in M.EVENT_ADDONS]
    lines += ["", "## Phụ phí", "", "| Khoản | Mức |", "|---|---|"]
    lines += [f"| {n} | {v} |" for n, v in M.EVENT_SURCHARGES]
    con = POLICY["consultation"]
    lines += ["", "## Thanh toán và hoàn cọc", "",
              f"- Đặt cọc **{con['contract']['deposit_percent']}%** khi ký hợp đồng; thanh toán phần còn lại trước "
              f"ngày sự kiện **{con['contract']['balance_due_days_before_event']} ngày**.",
              "- Đổi ngày sự kiện miễn phí **1 lần** nếu báo trước ít nhất **15 ngày**.",
              "- Hoàn cọc khi huỷ: ≥ 30 ngày hoàn 100%; 15–29 ngày hoàn 50%; dưới 15 ngày không hoàn. "
              "Mọi khoản hoàn cọc do quản lý studio duyệt.", "",
              "> Giá hoa nhập khẩu (mẫu đơn, tulip, hồng Ecuador) có thể thay đổi theo mùa; báo giá chính thức "
              "được gửi trong 48 giờ sau buổi tư vấn và giữ nguyên khi đã ký hợp đồng."]
    return "\n".join(lines)


def catalogue_2010_md() -> str:
    products = {p["sku"]: p for p in M.PRODUCTS}
    pick = ["BH02", "BH04", "BH07", "BH11", "BH06", "BH03", "GH01", "GH03", "GH02", "LC03", "HS01"]
    code = next(c for c in M.DISCOUNT_CODES if c["code"] == "PNVN2010")
    lines = ["# Catalogue hoa 20/10/2026",
             "Mừng ngày Phụ nữ Việt Nam – gửi lời thương đến mẹ, vợ, người yêu, cô giáo và đồng nghiệp.", "",
             "## Ưu đãi PNVN2010", "",
             f"- Giảm **{code['value']}%**, tối đa **{vnd(code['max_discount'])}** cho bó hoa, giỏ hoa, hộp hoa.",
             f"- Đơn tối thiểu **{vnd(code['min_order'])}** (tiền hoa, chưa gồm phí ship).",
             "- Hiệu lực theo ngày đặt đơn: **10/10/2026 – 20/10/2026**, mỗi khách 1 lần, không cộng dồn ưu đãi khác.",
             "", "## Mẫu gợi ý", "", "| Mã | Mẫu | Hoa chính | Giá |", "|---|---|---|---|"]
    lines += [f"| {s} | {products[s]['name']} | {products[s]['flowers']} | {vnd(products[s]['price'])} |" for s in pick]
    lines += ["", "## Quà tặng kèm", "", "| Mã | Sản phẩm | Giá |", "|---|---|---|"]
    lines += [f"| {s} | {products[s]['name']} | {vnd(products[s]['price'])} |" for s in ("PK01", "PK02", "PK03")]
    lines += ["", "## Lưu ý giao hàng ngày 20/10 (Thứ Ba)", "",
              "- 20/10 là **ngày cao điểm**: giờ giao có thể lệch ±30 phút so với khung đã chọn; **không nhận giao hỏa tốc**.",
              "- Nên đặt trước **18/10/2026** để giữ khung giờ mong muốn; khung K1–K3 thường kín sớm.",
              "- Phí ship: nội ô Ninh Kiều 25.000đ (miễn phí từ 500.000đ), Bình Thủy – Cái Răng 35.000đ (miễn phí từ 800.000đ).",
              "- Bó Tulip Hà Lan có màu theo lô nhập trong tuần; hoa tự nhiên có thể khác ảnh 10–15%.", "",
              "## Đặt hoa", "",
              "Website mocmienflorist.vn · Zalo 0939 868 186 · Hotline 0292 3 868 186 · Tại quầy Ninh Kiều và Cái Răng.", "",
              "> Thứ Tư Hồng: mỗi Thứ Tư tháng 10 – 11, giảm 15% hồng lẻ theo cành khi mua trực tiếp tại quầy "
              "(không áp dụng đơn online)."]
    return "\n".join(lines)


def build_pdfs() -> None:
    init_styles()
    render_pdf("22-bang-gia-hoa-cuoi-su-kien-2026.pdf", "Bảng giá hoa cưới & sự kiện 2026", wedding_price_md())
    src = (TOOLS / "pdf_src" / "23-dieu-khoan-hop-dong-hoa-su-kien.md").read_text(encoding="utf-8")
    render_pdf("23-dieu-khoan-hop-dong-hoa-su-kien.pdf", "Điều khoản hợp đồng hoa sự kiện", src)
    render_pdf("24-catalogue-hoa-20-10-2026.pdf", "Catalogue hoa 20/10/2026", catalogue_2010_md())


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    print("Sinh tài liệu vào", OUT.relative_to(ROOT))
    build_csvs()
    build_xlsx()
    build_pdfs()
